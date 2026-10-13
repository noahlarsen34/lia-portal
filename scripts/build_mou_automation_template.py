from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ORIGINAL_SOURCE = Path("/Users/noahlarsen/Downloads/New MOU Template.docx")
OUTPUT = Path("templates/mou/lia-mou-automation-template.docx")
SOURCE = ORIGINAL_SOURCE if ORIGINAL_SOURCE.exists() else OUTPUT

SCHOOL_PLACEHOLDER = "{{school_name}}"
PARTY_PLACEHOLDER = SCHOOL_PLACEHOLDER


def replace_in_runs(paragraph, old: str, new: str) -> int:
    replacements = 0

    for run in paragraph.runs:
        count = run.text.count(old)
        if count:
            run.text = run.text.replace(old, new)
            replacements += count

    return replacements


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))

    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)

    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")

    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)

    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        margin = margins.find(qn(f"w:{margin_name}"))
        if margin is None:
            margin = OxmlElement(f"w:{margin_name}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "D9D9D9")


def remove_table_borders(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def set_table_width(table, width_dxa: int, first_column_dxa: int = 7900) -> None:
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")

    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)

    table_width.set(qn("w:w"), str(width_dxa))
    table_width.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    grid_columns = grid.findall(qn("w:gridCol"))
    column_widths = (first_column_dxa, width_dxa - first_column_dxa)
    for column, width in zip(grid_columns, column_widths):
        column.set(qn("w:w"), str(width))


def set_cell_text(cell, text: str, *, bold=False, color="222222", align=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arimo"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Arimo")
    fonts.set(qn("w:hAnsi"), "Arimo")


def find_pricing_table(document: Document):
    for table in document.tables:
        table_text = " ".join(cell.text for row in table.rows for cell in row.cells)
        if "INITIAL START-UP COSTS" in table_text or "Initial Start-up Costs" in table_text:
            return table
    raise RuntimeError("The MOU pricing table could not be located.")


def build_pricing_table(document: Document) -> None:
    old_table = find_pricing_table(document)
    new_table = document.add_table(rows=27, cols=2)
    new_table.autofit = False
    set_table_width(new_table, 10725)
    set_table_borders(new_table)

    for row in new_table.rows:
        row.cells[0].width = Inches(5.5)
        row.cells[1].width = Inches(1.95)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=70, start=120, bottom=70, end=120)

        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)

    startup_header = new_table.cell(0, 0).merge(new_table.cell(0, 1))
    set_cell_shading(startup_header, "C4122F")
    set_cell_text(
        startup_header,
        "INITIAL START-UP COSTS - {{launch_school_year}}",
        bold=True,
        color="FFFFFF",
    )

    first_year_header = new_table.cell(6, 0).merge(new_table.cell(6, 1))
    set_cell_shading(first_year_header, "C4122F")
    set_cell_text(
        first_year_header,
        "LIA PROGRAM COSTS — FIRST YEAR - {{launch_school_year}}",
        bold=True,
        color="FFFFFF",
    )

    renewal_header = new_table.cell(17, 0).merge(new_table.cell(17, 1))
    set_cell_shading(renewal_header, "C4122F")
    set_cell_text(
        renewal_header,
        "ANNUAL RENEWAL PRICING - {{renewal_school_year}}",
        bold=True,
        color="FFFFFF",
    )

    rows = (
        (1, "Start-up fee, onboarding, and first-year training — per Middle or High School", "$2,500", False),
        (2, "Launching LIA: implementation training, teacher certification, and administrator/counselor training", "Included", False),
        (3, "Travel fees: Dr. José Enriquez initial meeting and founder visit", "$2,000", False),
        (4, "{{startup_discount_name}} ({{startup_discount_percent}}%)", "- {{startup_discount_amount}}", False),
        (5, "Start-up Fee — Total", "{{startup_final_price}}", True),
        (7, "LIA Portal Access", "$2,000", False),
        (8, "Access to Online Curriculum", "$1,000", False),
        (9, "Personalized Teacher Support throughout the school year — unlimited emails, calls, and meetings", "$1,400", False),
        (10, "LIA Youth Leadership Conference", "$800", False),
        (11, "LIA Leadership Boot Camp", "$500", False),
        (12, "Educator Institute", "$200", False),
        (13, "Masterminds", "$100", False),
        (14, "{{first_year_discount_name}} ({{first_year_discount_percent}}%)", "- {{first_year_discount_amount}}", False),
        (15, "First-Year Program Total", "{{first_year_final_price}}", True),
        (16, "Grand Total First-Year Implementation", "{{first_year_grand_total}}", True),
        (18, "LIA Portal Access", "$2,000", False),
        (19, "Access to Online Curriculum", "$1,000", False),
        (20, "Personalized Teacher Support throughout the school year — unlimited emails, calls, and meetings", "$1,400", False),
        (21, "LIA Youth Leadership Conference", "$800", False),
        (22, "LIA Leadership Boot Camp", "$500", False),
        (23, "Educator Institute", "$200", False),
        (24, "Masterminds", "$100", False),
        (25, "{{renewal_discount_name}} ({{renewal_discount_percent}}%)", "- {{renewal_discount_amount}}", False),
        (26, "Total Second-Year Annual Renewal Membership Fee", "{{renewal_program_final_price}}", True),
    )

    for row_index, label, amount, is_total in rows:
        if is_total:
            set_cell_shading(new_table.cell(row_index, 0), "F2F2F2")
            set_cell_shading(new_table.cell(row_index, 1), "F2F2F2")
        set_cell_text(new_table.cell(row_index, 0), label, bold=is_total)
        set_cell_text(
            new_table.cell(row_index, 1),
            amount,
            bold=is_total,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )

    old_table._tbl.addprevious(new_table._tbl)
    old_table._element.getparent().remove(old_table._element)

    first_paragraph = new_table._tbl.getprevious()
    if first_paragraph is not None and first_paragraph.tag == qn("w:p"):
        paragraph_properties = first_paragraph.get_or_add_pPr()
        page_break = OxmlElement("w:pageBreakBefore")
        paragraph_properties.append(page_break)


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def set_paragraph_font(paragraph, *, size=10, bold=False, color="222222") -> None:
    for run in paragraph.runs:
        run.font.name = "Arimo"
        run.font.size = Pt(size)
        run.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        run_properties = run._element.get_or_add_rPr()
        fonts = run_properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run_properties.insert(0, fonts)
        fonts.set(qn("w:ascii"), "Arimo")
        fonts.set(qn("w:hAnsi"), "Arimo")


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arimo"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string("1F1F1F")


def add_form_field(table, row_index: int, label: str, *, response_lines=1) -> None:
    label_cell = table.cell(row_index, 0)
    response_cell = table.cell(row_index, 1)
    set_cell_text(label_cell, label, bold=True)
    response_text = "\n".join("________________________________" for _ in range(response_lines))
    set_cell_text(response_cell, response_text)


def rebuild_post_pricing_sections(document: Document) -> None:
    pricing_table = find_pricing_table(document)
    body = document.element.body
    sibling = pricing_table._tbl.getnext()
    while sibling is not None and sibling.tag != qn("w:sectPr"):
        next_sibling = sibling.getnext()
        body.remove(sibling)
        sibling = next_sibling

    note = document.add_paragraph(
        "*After the second-year renewal, future renewal pricing will be based on "
        "the second-year renewal amount plus the applicable inflation percentage."
    )
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(0)
    set_paragraph_font(note, size=9)

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    add_section_heading(document, "AUTHORIZED SIGNATORIES")

    intro = document.add_paragraph(
        "The undersigned parties agree to the expectations and terms stated in this "
        "Memorandum of Understanding."
    )
    intro.paragraph_format.space_after = Pt(18)
    set_paragraph_font(intro, size=10)

    signature_table = document.add_table(rows=4, cols=2)
    signature_table.autofit = False
    set_table_width(signature_table, 10725, first_column_dxa=5362)
    for row in signature_table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=100, start=120, bottom=100, end=120)

    set_cell_text(signature_table.cell(0, 0), "LATINOS IN ACTION", bold=True, color="C4122F")
    set_cell_text(signature_table.cell(0, 1), "SCHOOL / DISTRICT", bold=True, color="C4122F")
    set_cell_text(signature_table.cell(1, 0), "Dr. José Enriquez\nCEO and Founder")
    set_cell_text(signature_table.cell(1, 1), "{{authorized_signer_name}}\nAuthorized Representative")

    signature_cell = signature_table.cell(2, 0)
    signature_cell.text = ""
    signature_paragraph = signature_cell.paragraphs[0]
    signature_paragraph.paragraph_format.space_after = Pt(0)
    signature_blob = None
    for relationship in document.part.rels.values():
        if "image" in relationship.reltype and relationship.target_ref.endswith("image1.png"):
            signature_blob = relationship.target_part.blob
            break
    if signature_blob:
        signature_paragraph.add_run().add_picture(BytesIO(signature_blob), width=Inches(2.1))
    else:
        signature_paragraph.add_run("________________________________")

    set_cell_text(signature_table.cell(2, 1), "________________________________\nSignature")
    set_cell_text(signature_table.cell(3, 0), "{{lia_signature_date}}\nDate")
    set_cell_text(signature_table.cell(3, 1), "________________________________\nDate")

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    add_section_heading(document, "INVOICE PROCESS")

    invoice_intro = document.add_paragraph(
        "Thank you for joining the Latinos In Action familia. After the MOU is signed, "
        "LIA will send an invoice for the agreed fees. Please provide the information "
        "below so the invoice can be processed correctly."
    )
    invoice_intro.paragraph_format.space_after = Pt(14)
    set_paragraph_font(invoice_intro, size=10)

    invoice_table = document.add_table(rows=10, cols=2)
    invoice_table.autofit = False
    set_table_width(invoice_table, 10725, first_column_dxa=2700)
    set_table_borders(invoice_table)
    for row in invoice_table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=120, start=140, bottom=120, end=140)

    sections = (
        (0, "PURCHASE ORDER", "Does your organization require a purchase order?  ☐ Yes  ☐ No"),
        (1, "PO contact", "Name of person who receives the quote"),
        (2, "PO contact email", "Email address"),
        (3, "PO contact phone", "Phone number"),
        (4, "Quote requirements", "Information that must appear on the quote"),
        (5, "INVOICE CONTACT", "Complete this section if no purchase order is required."),
        (6, "Invoice contact", "Name of person who receives the invoice"),
        (7, "Invoice email", "Email address"),
        (8, "Invoice phone", "Phone number"),
        (9, "VENDOR SETUP", "Is vendor registration required?  ☐ Yes  ☐ No\nRequired forms/documents:"),
    )
    for row_index, label, prompt in sections:
        if row_index in (0, 5, 9):
            set_cell_shading(invoice_table.cell(row_index, 0), "F2F2F2")
            set_cell_shading(invoice_table.cell(row_index, 1), "F2F2F2")
        set_cell_text(invoice_table.cell(row_index, 0), label, bold=True)
        set_cell_text(invoice_table.cell(row_index, 1), f"{prompt}\n________________________________________")

    return_note = document.add_paragraph(
        "Return vendor forms or invoicing questions to Brayan Zuniga at "
        "brayan@latinosinaction.org."
    )
    return_note.paragraph_format.space_before = Pt(12)
    return_note.paragraph_format.space_after = Pt(12)
    set_paragraph_font(return_note, size=10, bold=True, color="C4122F")

    closing = document.add_paragraph("Muchas gracias,\nLATINOS IN ACTION")
    set_paragraph_font(closing, size=10, bold=True)


def remove_legacy_page_footer(document: Document) -> None:
    """Remove the old PAGE / NUMPAGES field that some Word viewers show as only 'of'."""
    for section in document.sections:
        for footer in (
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            for paragraph in footer.paragraphs:
                for child in list(paragraph._p):
                    if child.tag != qn("w:pPr"):
                        paragraph._p.remove(child)


def rebuild_cover_header(document: Document) -> None:
    """Use balanced columns so Word and browser renderers share the same center."""
    for table in document.tables:
        table_text = " ".join(cell.text for row in table.rows for cell in row.cells)
        if "MEMORANDUM OF UNDERSTANDING" in table_text:
            return

    first_paragraphs = document.paragraphs[:5]
    if len(first_paragraphs) < 5:
        raise RuntimeError("The MOU cover header could not be located.")

    logo_blob = None
    for relationship in document.part.rels.values():
        if "image" in relationship.reltype and relationship.target_ref.endswith("image3.png"):
            logo_blob = relationship.target_part.blob
            break

    header_table = document.add_table(rows=1, cols=3)
    header_table.autofit = False
    set_table_width(header_table, 10725, first_column_dxa=1300)
    remove_table_borders(header_table)

    column_widths = (1300, 8125, 1300)
    grid_columns = header_table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_column, width in zip(grid_columns, column_widths):
        grid_column.set(qn("w:w"), str(width))

    for cell, width in zip(header_table.rows[0].cells, column_widths):
        cell.width = Inches(width / 1440)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell_properties = cell._tc.get_or_add_tcPr()
        cell_width = cell_properties.find(qn("w:tcW"))
        if cell_width is None:
            cell_width = OxmlElement("w:tcW")
            cell_properties.append(cell_width)
        cell_width.set(qn("w:w"), str(width))
        cell_width.set(qn("w:type"), "dxa")

    logo_cell = header_table.cell(0, 0)
    logo_cell.text = ""
    logo_paragraph = logo_cell.paragraphs[0]
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.space_before = Pt(0)
    logo_paragraph.paragraph_format.space_after = Pt(0)
    if logo_blob:
        logo_paragraph.add_run().add_picture(BytesIO(logo_blob), width=Inches(0.82))

    title_cell = header_table.cell(0, 1)
    title_cell.text = ""
    title_paragraph = title_cell.paragraphs[0]
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(2)
    title_run = title_paragraph.add_run("MEMORANDUM OF UNDERSTANDING")
    title_run.font.name = "Arimo"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string("000000")

    subtitle_paragraph = title_cell.add_paragraph()
    subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_paragraph.paragraph_format.space_before = Pt(0)
    subtitle_paragraph.paragraph_format.space_after = Pt(0)
    subtitle_paragraph.add_run("Between ")
    party_run = subtitle_paragraph.add_run(PARTY_PLACEHOLDER)
    party_run.bold = True
    subtitle_paragraph.add_run(" and Latinos In Action")
    set_paragraph_font(subtitle_paragraph, size=12)
    party_run.bold = True

    first_paragraphs[0]._p.addprevious(header_table._tbl)
    for paragraph in first_paragraphs:
        paragraph._element.getparent().remove(paragraph._element)


def normalize_responsibilities_and_implementation(document: Document) -> None:
    """Clean up the district-responsibility heading and implementation spacing."""
    paragraphs = document.paragraphs

    responsibilities_heading = paragraphs[60]
    replace_in_runs(
        responsibilities_heading,
        PARTY_PLACEHOLDER,
        "SCHOOL / DISTRICT RESPONSIBILITIES",
    )
    responsibilities_heading.paragraph_format.left_indent = Inches(0)
    responsibilities_heading.paragraph_format.space_before = Pt(8)
    responsibilities_heading.paragraph_format.space_after = Pt(6)
    responsibilities_heading.paragraph_format.keep_with_next = True
    set_paragraph_font(responsibilities_heading, size=14, bold=True)

    responsibilities_intro = paragraphs[61]
    responsibilities_intro.paragraph_format.space_before = Pt(0)
    responsibilities_intro.paragraph_format.space_after = Pt(8)

    school_name = paragraphs[74]
    school_name.paragraph_format.space_before = Pt(0)
    school_name.paragraph_format.space_after = Pt(4)
    school_name.paragraph_format.keep_with_next = True

    blank_after_school = paragraphs[75]
    if not blank_after_school.text.strip():
        blank_after_school._element.getparent().remove(blank_after_school._element)

    implementation_followup = document.paragraphs[75]
    implementation_followup.paragraph_format.space_before = Pt(0)
    implementation_followup.paragraph_format.space_after = Pt(12)


def build_template() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source template not found: {SOURCE}")

    document = Document(SOURCE)
    original_placeholder = "[Name of School or District]"
    party_replacements = 0
    school_replacements = 0

    for paragraph_index, paragraph in enumerate(document.paragraphs):
        if paragraph_index == 74:
            school_replacements += replace_in_runs(
                paragraph,
                original_placeholder,
                SCHOOL_PLACEHOLDER,
            )
        else:
            party_replacements += replace_in_runs(
                paragraph,
                original_placeholder,
                PARTY_PLACEHOLDER,
            )

        replace_in_runs(paragraph, "2025-2026", "{{effective_school_year}}")
        replace_in_runs(paragraph, "03/24/26", "{{lia_signature_date}}")
        replace_in_runs(
            paragraph,
            f"{PARTY_PLACEHOLDER}and LIA",
            f"{PARTY_PLACEHOLDER} and LIA",
        )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    party_replacements += replace_in_runs(
                        paragraph,
                        original_placeholder,
                        PARTY_PLACEHOLDER,
                    )

    is_existing_automation_template = SOURCE == OUTPUT

    if not is_existing_automation_template and (
        party_replacements != 23 or school_replacements != 1
    ):
        raise RuntimeError(
            "Unexpected school/district placeholder count: "
            f"party={party_replacements}, school={school_replacements}"
        )

    normalize_responsibilities_and_implementation(document)
    build_pricing_table(document)
    rebuild_post_pricing_sections(document)
    remove_legacy_page_footer(document)
    rebuild_cover_header(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_template()
